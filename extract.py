import os
try:
    from pptx import Presentation
except ImportError:
    os.system('pip install python-pptx -q')
    from pptx import Presentation

prs = Presentation('ProjectSummary.pptx')
with open('extracted_pptx.txt', 'w', encoding='utf-8') as f:
    for i, slide in enumerate(prs.slides):
        f.write(f'--- Slide {i+1} ---\n')
        for shape in slide.shapes:
            if hasattr(shape, 'text'):
                f.write(shape.text + '\n')
